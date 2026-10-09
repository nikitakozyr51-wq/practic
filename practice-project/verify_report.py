"""Check source integrity, equipment photo attribution and report structure."""
from pathlib import Path
from zipfile import ZipFile
from hashlib import sha256
from lxml import etree
import json,re
from docx import Document

root=Path(__file__).resolve().parent
a=root/'analysis'
s=(a/'REPORT_CONTENT.md').read_text()
content=s[s.index('## 1.'):]
main=[int(x) for x in re.findall(r'^## (\d+)\.',content,re.M)]
subs=re.findall(r'^### (\d+\.\d+)\.',content,re.M)
assert main==list(range(1,18)),main
expected=[f'{n}.{i}' for n,count in [(1,3),(2,4),(3,3),(4,4),(5,5),(6,3),(7,4),(8,7)] for i in range(1,count+1)]
assert subs==expected,subs
body,bib=content.split('## Список использованной литературы',1)
cites=set(re.findall(r'\[(\d+)\]',body))
refs=set(re.findall(r'^(\d+)\.',bib,re.M))
assert cites==refs=={str(i) for i in range(1,8)},(cites,refs)
assert 'Шеньшин' not in content
assert '62 000 кВт' in body and '6200' not in body
assert '7570' in body and '5200 кВт' in body and 'GTA M-9' in body
figures=json.loads((a/'report-figures.json').read_text())
assert {int(n) for n in re.findall(r'рисунке (\d+)',body)}=={p['number'] for p in figures}
assert [int(n) for n in re.findall(r'<!-- figure:(\d+) -->',body)]==[p['number'] for p in figures]
report=a/'report.docx'
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships','a':'http://schemas.openxmlformats.org/drawingml/2006/main','wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}
receipt=json.loads((a/'report-build-receipt.json').read_text())
with ZipFile(report) as z,ZipFile(root/'originals'/'Nazar_Otchet_Praktika.docx') as original,ZipFile(root/'originals'/'ОТЧЕТ.docx') as maxim:
    assert z.testzip() is None
    tree=etree.fromstring(z.read('word/document.xml'))
    source_tree=etree.fromstring(original.read('word/document.xml'))
    assert etree.tostring(tree.find('.//w:sectPr',ns))==etree.tostring(source_tree.find('.//w:sectPr',ns))
    for name in ['word/styles.xml','word/theme/theme1.xml','word/numbering.xml','word/footer1.xml','word/footer2.xml']:
        assert z.read(name)==original.read(name),name
    for p in receipt['photos']:
        data=z.read('word/media/equipment-figure-'+str(p['number'])+'.jpeg')
        if p['source_kind']=='provided_docx':
            assert data==maxim.read('word/media/'+p['original_media']),p
        elif p['source_kind']=='manufacturer':
            assert data==(root/p['source_path']).read_bytes(),p
            assert sha256(data).hexdigest()==p['source_sha256'],p
        else:raise AssertionError(p)
        assert sha256(data).hexdigest()==p['sha256']
    ids=tree.xpath('.//wp:docPr/@id',namespaces=ns)
    assert len(ids)==len(set(ids)),ids
    rels=etree.fromstring(z.read('word/_rels/document.xml.rels'))
    by_id={r.get('Id'):r.get('Target') for r in rels}
    for rid in tree.xpath('.//a:blip/@r:embed',namespaces=ns):
        assert 'word/'+by_id[rid] in z.namelist(),rid
    texts=tree.xpath('.//w:t/text()',namespaces=ns)
    full=' '.join(texts)
    assert 'PANTHERLIGHT' not in full and 'Еремян' not in full
    assert 'Козырева Никиты Алексеевича' in full and '[ФИО]' not in full
    assert sum(t.startswith('Рис. ') for t in texts)==len(figures)
    assert 'Приложение А. Задание на производственную практику' in full
    assert 'Шеньшин' not in full and 'Максим' not in full
    assert '<!-- figure:' not in full and '####' not in full
    assert receipt['assignment']['native_word_text']
    assert receipt['assignment']['all_source_wording_preserved']
    core=z.read('docProps/core.xml').decode()
    assert 'Назар' not in core and 'Козырев Никита Алексеевич' in core

doc=Document(report)
first=next(i for i,p in enumerate(doc.paragraphs)
           if p.text==receipt['paragraph_formats'][0]['text'])
for record,p in zip(receipt['paragraph_formats'],doc.paragraphs[first:]):
    assert p.text==record['text'],(record,p.text)
    role=record['role']
    assert p.paragraph_format.line_spacing==1,(role,p.text[:80])
    if role in ('body','heading','device_heading'):
        assert str(p.alignment)=='JUSTIFY (3)',(role,p.text[:80])
        assert abs(p.paragraph_format.first_line_indent.cm-1.25)<0.002
    elif role in ('caption','image','appendix','bibliography'):
        assert str(p.alignment)=='CENTER (1)',(role,p.text[:80])
    if role in ('body','device_heading') or (role=='heading' and re.match(r'^\d+\.',p.text)):
        assert p._p.find('w:pPr/w:numPr',p._p.nsmap) is None,p.text[:80]
    if role in ('heading','device_heading','image','bibliography','appendix'):
        assert p.paragraph_format.keep_with_next is True,p.text[:80]
    for run in p.runs:
        if run.text:
            assert run.font.name=='Times New Roman' and run.font.size.pt==12,p.text[:80]
            assert str(run.font.color.rgb)=='000000',p.text[:80]
            if role=='device_heading':assert run.bold is True,p.text
page_map=json.loads((a/'report-page-map.json').read_text())
for i,p in enumerate(doc.paragraphs[16:33],1):
    fields=p.text.rsplit('\t',1)
    assert len(fields)==2 and fields[1]==str(page_map[str(i)]),(i,p.text)
    assert not re.search(r'\d+$',fields[0]),('old page number in heading',p.text)
rows=[]
split=re.split(r'^(#{2,4}) ([^\n]+)\n',body,flags=re.M)
for level,title,prose in zip(split[1::3],split[2::3],split[3::3]):
    paragraphs=[p.strip() for p in prose.strip().split('\n\n')
                if p.strip() and not p.strip().startswith('<!-- figure:')]
    for i,p in enumerate(paragraphs,1):
        rows.append({'section':title,'paragraph':i,'text':p,'citations':re.findall(r'\[(\d+)\]',p),'review_scope':'Термины, единицы, физический смысл, согласованность; наличие оборудования и результаты наблюдений не верифицированы.'})
(a/'thesis-register.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2))
result={'main_sections':len(main),'subsections':len(subs),'assignment_leaf_topics':len(subs)+9,'body_paragraphs':len(rows),'source_photos':len(receipt['photos']),'source_assignment_paragraphs':receipt['assignment']['source_paragraphs'],'assignment_native_word_paragraphs':receipt['assignment']['inserted_nonempty_paragraphs'],'citations_and_references':sorted(cites),'orphans':0,'peer_report_removed_from_bibliography':True,'original_photo_bytes_preserved':True,'template_parts_and_section_settings_identical':True,'report_sha256':sha256(report.read_bytes()).hexdigest(),'report_bytes':report.stat().st_size}
(a/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
print(json.dumps(result,ensure_ascii=False))
